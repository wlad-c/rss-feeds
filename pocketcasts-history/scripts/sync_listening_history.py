#!/usr/bin/env python3
"""Fetch latest Pocket Casts listening history and rebuild Listening Log tables.

Requires env:
  POCKETCAST_EMAIL
  POCKETCAST_PASSWORD

Writes under pocketcasts-history/tables/, refreshes Excel + standalone HTML,
syncs the slim docs/ runtime copy, and remaps notes/themes for new episodes.
"""

from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
TABLES = ROOT / "tables"
DOCS = REPO / "docs"
YEARS_DIR = TABLES / "years"

STATUS = {0: "Unplayed", 1: "In progress", 2: "In progress", 3: "Played"}
API = "https://api.pocketcasts.com"
UA = "PocketCasts/7.0"


def die(msg: str, code: int = 1) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    raise SystemExit(code)


def post(url: str, body: dict | None = None, token: str | None = None) -> dict:
    data = None if body is None else json.dumps(body).encode()
    headers = {
        "Content-Type": "application/json",
        "Accept": "application/json",
        "User-Agent": UA,
    }
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, data=data or b"{}", headers=headers, method="POST")
    with urllib.request.urlopen(req, timeout=120) as resp:
        return json.loads(resp.read().decode())


def fmt_secs(seconds: int | float | str | None) -> str:
    if seconds in (None, ""):
        return ""
    try:
        s = int(float(seconds))
    except (TypeError, ValueError):
        return ""
    s = max(0, s)
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m}:{sec:02d}"


def parse_hms_to_sec(value: int | float | str | None) -> int | None:
    """Parse seconds int or m:ss / h:mm:ss string. Returns None when empty/invalid."""
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            s = int(value)
        except (TypeError, ValueError):
            return None
        return s if s > 0 else None
    text = str(value).strip()
    if not text:
        return None
    if text.isdigit():
        s = int(text)
        return s if s > 0 else None
    parts = text.split(":")
    try:
        nums = [int(p) for p in parts]
    except ValueError:
        return None
    if len(nums) == 3:
        s = nums[0] * 3600 + nums[1] * 60 + nums[2]
    elif len(nums) == 2:
        s = nums[0] * 60 + nums[1]
    elif len(nums) == 1:
        s = nums[0]
    else:
        return None
    return s if s > 0 else None


def resolve_duration_sec(
    dur_sec: int | str | None, played_sec: int | str | None
) -> tuple[int | str, bool]:
    """Prefer catalog duration; fall back to played-up-to when duration is missing/0.

    Pocket Casts often returns duration=0 while playedUpTo still has real listen
    time. Hours charts use duration, so without this fallback Played episodes
    count in Eps mode but show 0h.
    """
    d = parse_hms_to_sec(dur_sec)
    if d is not None:
        return d, False
    p = parse_hms_to_sec(played_sec)
    if p is not None:
        return p, True
    # Preserve empty vs explicit 0 from the API when nothing is recoverable.
    if dur_sec == 0 or dur_sec == "0":
        return 0, False
    return "", False


def ms_to_dt(ms: int | str) -> datetime:
    value = int(ms)
    if value > 1e12:
        value //= 1000
    return datetime.fromtimestamp(value, timezone.utc)


def login() -> str:
    email = os.environ.get("POCKETCAST_EMAIL", "").strip()
    password = os.environ.get("POCKETCAST_PASSWORD", "").strip()
    if not email or not password:
        die("POCKETCAST_EMAIL and POCKETCAST_PASSWORD must be set")
    data = post(
        f"{API}/user/login",
        {"email": email, "password": password, "scope": "webplayer"},
    )
    token = data.get("token")
    if not token:
        die("login succeeded but no token returned")
    print(f"logged in as {data.get('email') or email}")
    return token


def fetch_catalog(token: str) -> tuple[dict[str, dict], dict[str, str]]:
    plist = post(f"{API}/user/podcast/list", {"v": 1}, token)
    podcasts = {p["uuid"]: p for p in (plist.get("podcasts") or []) if p.get("uuid")}
    folders = {
        f["folderUuid"]: (f.get("name") or "")
        for f in (plist.get("folders") or [])
        if f.get("folderUuid")
    }
    print(f"podcasts={len(podcasts)} folders={len(folders)}")
    return podcasts, folders


def fetch_progress(token: str, podcasts: dict[str, dict]) -> dict[str, dict]:
    progress: dict[str, dict] = {}
    for i, uuid in enumerate(podcasts, 1):
        ep = post(f"{API}/user/podcast/episodes", {"uuid": uuid}, token)
        for row in ep.get("episodes") or []:
            eu = row.get("uuid")
            if eu:
                progress[eu] = row
        if i % 50 == 0 or i == len(podcasts):
            print(f"  progress {i}/{len(podcasts)} ({len(progress)} rows)")
    return progress


def fetch_history(token: str, start_year: int = 2010) -> tuple[list[dict], dict]:
    now_year = datetime.now(timezone.utc).year
    year_counts: dict[str, int] = {}
    changes: list[dict] = []
    for year in range(start_year, now_year + 1):
        resp = post(
            f"{API}/history/year",
            {"year": year, "version": "1", "count": False},
            token,
        )
        rows = ((resp.get("history") or {}).get("changes")) or []
        year_counts[str(year)] = len(rows)
        print(f"  year {year}: {len(rows)}")
        for row in rows:
            row["_year"] = year
            changes.append(row)
    return changes, year_counts


def fetch_stats(token: str) -> dict:
    return post(f"{API}/user/stats/summary", {}, token)


def fetch_starred(token: str) -> set[str]:
    try:
        data = post(f"{API}/user/starred", {}, token)
    except Exception as exc:  # noqa: BLE001
        print(f"starred fetch failed: {exc}")
        return set()
    return {e.get("uuid") for e in (data.get("episodes") or []) if e.get("uuid")}


def build_rows(
    changes: list[dict],
    podcasts: dict[str, dict],
    folders: dict[str, str],
    progress: dict[str, dict],
    starred: set[str],
) -> list[dict]:
    rows: list[dict] = []
    for c in changes:
        ep_uuid = c.get("episode") or ""
        pod_uuid = c.get("podcast") or ""
        p = podcasts.get(pod_uuid) or {}
        prog = progress.get(ep_uuid) or {}
        listened = ms_to_dt(c["modifiedAt"])
        pub = c.get("published") or ""
        duration = prog.get("duration")
        played_up = prog.get("playedUpTo")
        status_code = prog.get("playingStatus")
        raw_dur: int | str
        played_sec: int | str
        try:
            raw_dur = int(duration) if duration not in (None, "") else ""
        except (TypeError, ValueError):
            raw_dur = ""
        try:
            played_sec = int(played_up) if played_up not in (None, "") else ""
        except (TypeError, ValueError):
            played_sec = ""
        dur_sec, _from_played = resolve_duration_sec(raw_dur, played_sec)
        prog_pct = ""
        if dur_sec not in ("", 0) and played_sec != "":
            prog_pct = f"{min(100.0, round(100.0 * float(played_sec) / float(dur_sec), 1))}"
        status = STATUS.get(status_code, "") or "Played"
        folder_uuid = p.get("folderUuid")
        rows.append(
            {
                "Listened at (UTC)": listened.strftime("%Y-%m-%d %H:%M:%S"),
                "Listened date": listened.strftime("%Y-%m-%d"),
                "Year": listened.year,
                "Podcast": p.get("title") or "",
                "Episode": c.get("title") or "",
                "Author": p.get("author") or "",
                "Folder": folders.get(folder_uuid, "") if folder_uuid else "",
                "Published": pub[:10] if pub else "",
                "Duration": fmt_secs(dur_sec) if dur_sec != "" else "",
                "Duration (sec)": dur_sec if dur_sec != "" else "",
                "Played up to": fmt_secs(played_sec) if played_sec != "" else "",
                "Progress %": prog_pct,
                "Status": status,
                "Starred": "Yes"
                if (prog.get("starred") or ep_uuid in starred)
                else "",
                "Still subscribed": "Yes" if pod_uuid in podcasts else "No",
                "Episode URL": c.get("url") or "",
                "Podcast feed": p.get("webFeed")
                if isinstance(p.get("webFeed"), str)
                else "",
                "Episode UUID": ep_uuid,
                "Podcast UUID": pod_uuid,
            }
        )
    rows.sort(key=lambda r: r["Listened at (UTC)"], reverse=True)
    return rows


def to_compact(r: dict) -> dict:
    return {
        "t": r["Listened at (UTC)"],
        "d": r["Listened date"],
        "y": int(r["Year"]),
        "p": r.get("Podcast") or "",
        "e": r.get("Episode") or "",
        "a": r.get("Author") or "",
        "f": r.get("Folder") or "",
        "pub": r.get("Published") or "",
        "dur": r.get("Duration") or "",
        "prog": r.get("Progress %") or "",
        "st": r.get("Status") or "",
        "star": r.get("Starred") == "Yes",
    }


def parse_dur_sec(r: dict) -> int | None:
    v = r.get("Duration (sec)")
    parsed = parse_hms_to_sec(v)
    if parsed is not None:
        return parsed
    return parse_hms_to_sec(r.get("Duration") or "")


def repair_missing_durations(rows: list[dict]) -> int:
    """Fill Duration from Played up to when catalog duration is missing/zero.

    Returns the number of rows repaired. Safe to run repeatedly.
    """
    fixed = 0
    for r in rows:
        existing = parse_hms_to_sec(r.get("Duration (sec)")) or parse_hms_to_sec(
            r.get("Duration")
        )
        if existing is not None:
            continue
        played_sec = parse_hms_to_sec(r.get("Played up to"))
        if played_sec is None:
            continue
        r["Duration (sec)"] = played_sec
        r["Duration"] = fmt_secs(played_sec)
        # played-up-to is the only known length, so progress is complete vs that.
        if r.get("Progress %") in (None, ""):
            r["Progress %"] = "100.0"
        fixed += 1
    return fixed


def build_aggregates(rows: list[dict], stats: dict, year_counts_raw: dict) -> tuple[dict, dict]:
    pod_stats: dict[str, dict] = defaultdict(
        lambda: {
            "n": 0,
            "first": None,
            "last": None,
            "a": "",
            "sub": False,
            "feed": "",
            "uuid": "",
            "website": "",
        }
    )
    for r in rows:
        key = r.get("Podcast") or r["Podcast UUID"] or "Unknown"
        s = pod_stats[key]
        s["n"] += 1
        s["a"] = r.get("Author") or s["a"]
        s["sub"] = r.get("Still subscribed") == "Yes"
        s["feed"] = r.get("Podcast feed") or s["feed"]
        s["uuid"] = r.get("Podcast UUID") or s["uuid"]
        d = r["Listened date"]
        s["first"] = d if not s["first"] or d < s["first"] else s["first"]
        s["last"] = d if not s["last"] or d > s["last"] else s["last"]

    year_counts = Counter(int(r["Year"]) for r in rows)
    years_sorted = [str(y) for y in sorted(year_counts, reverse=True)]

    progress_vals: list[float] = []
    dur_secs: list[int] = []
    starred = 0
    monthly: Counter[str] = Counter()
    daily: Counter[str] = Counter()
    weekday: Counter[str] = Counter()
    hour: Counter[int] = Counter()
    folder_c: Counter[str] = Counter()
    status_c: Counter[str] = Counter()
    show_c: Counter[str] = Counter()
    dur_by_year: dict[int, list[int]] = defaultdict(list)
    wd_names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

    for r in rows:
        ds = parse_dur_sec(r)
        if ds is not None:
            dur_secs.append(ds)
            dur_by_year[int(r["Year"])].append(ds)
        try:
            if r.get("Progress %") not in (None, ""):
                progress_vals.append(float(r["Progress %"]))
        except ValueError:
            pass
        if r.get("Starred") == "Yes":
            starred += 1
        monthly[r["Listened date"][:7]] += 1
        daily[r["Listened date"]] += 1
        dt = datetime.strptime(r["Listened at (UTC)"], "%Y-%m-%d %H:%M:%S")
        weekday[wd_names[dt.weekday()]] += 1
        hour[dt.hour] += 1
        folder_c[r.get("Folder") or "Ungrouped"] += 1
        status_c[r.get("Status") or "Unknown"] += 1
        show_c[r.get("Podcast") or "Unknown"] += 1

    months_sorted = sorted(monthly)
    rolling = []
    for i, m in enumerate(months_sorted):
        window = months_sorted[max(0, i - 2) : i + 1]
        rolling.append(round(sum(monthly[x] for x in window) / len(window), 1))
    cum = []
    total = 0
    for m in months_sorted:
        total += monthly[m]
        cum.append({"month": m, "total": total})

    busiest_month = max(monthly.items(), key=lambda x: x[1]) if monthly else ("", 0)
    busiest_day = max(daily.items(), key=lambda x: x[1]) if daily else ("", 0)
    top_shows = show_c.most_common(40)
    top5 = sum(c for _, c in top_shows[:5])
    top1 = top_shows[0][1] if top_shows else 0

    latest = max(datetime.strptime(r["Listened date"], "%Y-%m-%d") for r in rows)
    start = latest - timedelta(days=365)
    heatmap = []
    d = start
    while d <= latest:
        key = d.strftime("%Y-%m-%d")
        heatmap.append({"date": key, "count": daily.get(key, 0), "dow": d.weekday()})
        d += timedelta(days=1)

    yoy: dict[str, dict[str, int]] = defaultdict(dict)
    for m, c in monthly.items():
        y, mo = m.split("-")
        yoy[mo][y] = c

    avg_dur_year = []
    for y in sorted(dur_by_year):
        vals = dur_by_year[y]
        avg_dur_year.append(
            {
                "year": y,
                "avgMinutes": round(sum(vals) / len(vals) / 60, 1),
                "episodes": len(vals),
            }
        )

    earliest = min(r["Listened at (UTC)"] for r in rows) if rows else None
    latest_listen = max(r["Listened at (UTC)"] for r in rows) if rows else None
    coverage = {
        "accountStartedAt": stats.get("timesStartedAt"),
        "timeListenedSeconds": int(stats.get("timeListened") or 0),
        "timeListenedHours": round(int(stats.get("timeListened") or 0) / 3600, 1),
        "timeSilenceRemovalSeconds": int(stats.get("timeSilenceRemoval") or 0),
        "timeSkippingSeconds": int(stats.get("timeSkipping") or 0),
        "timeIntroSkippingSeconds": int(stats.get("timeIntroSkipping") or 0),
        "timeVariableSpeedSeconds": int(stats.get("timeVariableSpeed") or 0),
        "historyEarliestListen": earliest,
        "historyLatestListen": latest_listen,
        "perEpisodeHistoryAvailableFrom": "2019-03-07",
        "yearsProbed": year_counts_raw,
        "note": (
            "Pocket Casts /history/year returns 0 episodes for years before 2019 "
            "for this account. Per-episode listening history begins 2019-03-07."
        ),
        "syncedAt": datetime.now(timezone.utc).isoformat(),
    }

    meta = {
        "title": "Pocket Casts listening history",
        "exportedAt": datetime.now(timezone.utc).isoformat(),
        "totalEpisodes": len(rows),
        "totalPodcasts": len(pod_stats),
        "years": years_sorted,
        "yearCounts": {str(y): year_counts[y] for y in sorted(year_counts, reverse=True)},
        "starredEpisodes": starred,
        "avgProgressPct": round(sum(progress_vals) / len(progress_vals), 1)
        if progress_vals
        else None,
        "totalDurationHours": round(sum(dur_secs) / 3600, 1) if dur_secs else 0,
        "avgEpisodeMinutes": round(sum(dur_secs) / len(dur_secs) / 60, 1)
        if dur_secs
        else None,
        "busiestMonth": {"month": busiest_month[0], "count": busiest_month[1]},
        "busiestDay": {"date": busiest_day[0], "count": busiest_day[1]},
        "topShowSharePct": round(100 * top1 / len(rows), 1) if rows else 0,
        "top5SharePct": round(100 * top5 / len(rows), 1) if rows else 0,
        "accountStartedAt": stats.get("timesStartedAt"),
        "accountTimeListenedHours": coverage["timeListenedHours"],
        "perEpisodeHistoryFrom": coverage["perEpisodeHistoryAvailableFrom"],
        "yearsWithoutHistory": [y for y, c in year_counts_raw.items() if c == 0],
    }

    aggregates = {
        "meta": meta,
        "monthly": [{"month": m, "count": monthly[m]} for m in months_sorted],
        "monthlyRolling3": rolling,
        "cumulative": cum,
        "yearly": [{"year": y, "count": year_counts[y]} for y in sorted(year_counts)],
        "weekday": [{"day": d, "count": weekday[d]} for d in wd_names],
        "hour": [{"hour": h, "count": hour[h]} for h in range(24)],
        "folders": [{"name": n, "count": c} for n, c in folder_c.most_common()],
        "status": [{"name": n, "count": c} for n, c in status_c.most_common()],
        "topPodcasts": [{"name": n, "count": c} for n, c in top_shows],
        "avgDurationByYear": avg_dur_year,
        "yoyByMonth": dict(yoy),
        "heatmap": heatmap,
        "coverage": coverage,
    }
    return {
        "meta": meta,
        "aggregates": aggregates,
        "coverage": coverage,
        "pod_stats": pod_stats,
        "year_counts": year_counts,
    }, coverage


def remap_notes(old_compact: list[dict], new_compact: list[dict]) -> dict:
    """Keep notes attached to episodes across re-exports (index shifts)."""
    notes_path = TABLES / "notes.json"
    if not notes_path.exists():
        return {}
    old_notes = json.loads(notes_path.read_text(encoding="utf-8"))
    if not isinstance(old_notes, dict) or not old_notes:
        return {}

    def key_for(ep: dict) -> tuple[str, str]:
        return (
            " ".join((ep.get("p") or "").casefold().split()),
            " ".join((ep.get("e") or "").casefold().split()),
        )

    old_by_key: dict[tuple[str, str], dict] = {}
    for idx_s, note in old_notes.items():
        try:
            idx = int(idx_s)
        except ValueError:
            continue
        if 0 <= idx < len(old_compact):
            old_by_key[key_for(old_compact[idx])] = note
        elif isinstance(note, dict) and (note.get("p") or note.get("e")):
            old_by_key[key_for({"p": note.get("p"), "e": note.get("e")})] = note

    remapped: dict[str, dict] = {}
    for i, ep in enumerate(new_compact):
        note = old_by_key.get(key_for(ep))
        if not note:
            continue
        entry = dict(note)
        entry["p"] = ep.get("p")
        entry["e"] = ep.get("e")
        remapped[str(i)] = entry
    print(f"remapped notes: {len(old_notes)} → {len(remapped)}")
    return remapped


def write_tables(rows: list[dict], built: dict) -> list[dict]:
    TABLES.mkdir(parents=True, exist_ok=True)
    YEARS_DIR.mkdir(parents=True, exist_ok=True)

    compact = [to_compact(r) for r in rows]
    fields = list(rows[0].keys()) if rows else []

    with (TABLES / "episodes.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    (TABLES / "episodes.json").write_text(
        json.dumps(compact, separators=(",", ":")), encoding="utf-8"
    )

    meta = built["meta"]
    (TABLES / "meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    (TABLES / "aggregates.json").write_text(
        json.dumps(built["aggregates"], separators=(",", ":")), encoding="utf-8"
    )
    (TABLES / "coverage.json").write_text(
        json.dumps(built["coverage"], indent=2), encoding="utf-8"
    )

    pod_stats = built["pod_stats"]
    podcasts_compact = [
        {
            "p": name,
            "a": s["a"],
            "n": s["n"],
            "first": s["first"],
            "last": s["last"],
            "sub": s["sub"],
        }
        for name, s in sorted(pod_stats.items(), key=lambda x: -x[1]["n"])
    ]
    (TABLES / "podcasts.json").write_text(
        json.dumps(podcasts_compact, separators=(",", ":")), encoding="utf-8"
    )
    with (TABLES / "podcasts.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "Podcast",
                "Author",
                "Episodes listened",
                "First listened",
                "Last listened",
                "Still subscribed",
                "Website",
                "RSS feed",
                "Podcast UUID",
            ],
        )
        w.writeheader()
        for name, s in sorted(pod_stats.items(), key=lambda x: -x[1]["n"]):
            w.writerow(
                {
                    "Podcast": name,
                    "Author": s["a"],
                    "Episodes listened": s["n"],
                    "First listened": s["first"],
                    "Last listened": s["last"],
                    "Still subscribed": "Yes" if s["sub"] else "No",
                    "Website": s.get("website") or "",
                    "RSS feed": s["feed"],
                    "Podcast UUID": s["uuid"],
                }
            )

    year_counts = built["year_counts"]
    with (TABLES / "years.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Year", "Episodes listened", "Distinct podcasts"])
        for y in sorted(year_counts):
            pods = len({r.get("Podcast") for r in rows if r["Year"] == y})
            w.writerow([y, year_counts[y], pods])

    for path in YEARS_DIR.glob("*.json"):
        path.unlink()
    for y in sorted(year_counts):
        yrows = [to_compact(r) for r in rows if r["Year"] == y]
        (YEARS_DIR / f"{y}.json").write_text(
            json.dumps(yrows, separators=(",", ":")), encoding="utf-8"
        )

    recent = rows[:500]
    with (TABLES / "recent.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(recent)
    (TABLES / "recent.json").write_text(
        json.dumps([to_compact(r) for r in recent], separators=(",", ":")),
        encoding="utf-8",
    )
    return compact


def write_excel(rows: list[dict], coverage: dict, meta: dict) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary = [
        ("Title", meta["title"]),
        ("Exported at", meta["exportedAt"]),
        ("Total episodes", meta["totalEpisodes"]),
        ("Total podcasts", meta["totalPodcasts"]),
        ("Account started", meta.get("accountStartedAt")),
        ("Account time listened (hours)", meta.get("accountTimeListenedHours")),
        ("Per-episode history from", meta.get("perEpisodeHistoryFrom")),
        ("Earliest listen", coverage.get("historyEarliestListen")),
        ("Latest listen", coverage.get("historyLatestListen")),
        ("Synced at", coverage.get("syncedAt")),
    ]
    for i, (k, v) in enumerate(summary, 1):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)

    fields = list(rows[0].keys()) if rows else []
    ws2 = wb.create_sheet("Episodes")
    ws2.append(fields)
    for r in rows:
        ws2.append([r.get(c, "") for c in fields])

    ws3 = wb.create_sheet("Recent 500")
    ws3.append(fields)
    for r in rows[:500]:
        ws3.append([r.get(c, "") for c in fields])

    ws4 = wb.create_sheet("Podcasts")
    with (TABLES / "podcasts.csv").open(encoding="utf-8") as f:
        for row in csv.reader(f):
            ws4.append(row)

    ws5 = wb.create_sheet("Years")
    with (TABLES / "years.csv").open(encoding="utf-8") as f:
        for row in csv.reader(f):
            ws5.append(row)

    ws6 = wb.create_sheet("Coverage")
    ws6.append(["Key", "Value"])
    for k, v in coverage.items():
        ws6.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v])

    wb.save(ROOT / "PocketCasts_Listening_History.xlsx")
    print("wrote Excel workbook")


def sync_docs() -> None:
    if not DOCS.exists():
        print("docs/ missing — skip Pages sync")
        return
    dest = DOCS / "tables"
    dest.mkdir(parents=True, exist_ok=True)
    for name in (
        "aggregates.json",
        "episodes.json",
        "podcasts.json",
        "themes.json",
        "episode-themes.json",
        "notes.json",
        "notes-meta.json",
    ):
        src = TABLES / name
        if src.exists():
            (dest / name).write_bytes(src.read_bytes())
    for name in ("explore.js", "explore.css", "index.html", "README.md"):
        src = ROOT / name
        if name == "README.md":
            # Keep docs README (Pages-specific); only copy explorer assets.
            continue
        if src.exists():
            (DOCS / name).write_bytes(src.read_bytes())
    print("synced slim tables into docs/")


def run_themes() -> None:
    script = ROOT / "scripts" / "rebuild_theme_hierarchy.py"
    if not script.exists():
        print("theme hierarchy script missing — skip")
        return
    print("rebuilding theme hierarchy…")
    subprocess.check_call([sys.executable, str(script)], cwd=str(ROOT))


def run_standalone() -> None:
    script = ROOT / "build-standalone.py"
    print("rebuilding standalone HTML…")
    subprocess.check_call([sys.executable, str(script)], cwd=str(ROOT))


def load_episode_csv() -> list[dict]:
    path = TABLES / "episodes.csv"
    if not path.exists():
        die(f"missing {path}")
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def stats_from_coverage(coverage: dict) -> dict:
    """Rebuild Pocket Casts stats-shaped dict from a prior coverage.json."""
    listened = int(coverage.get("timeListenedSeconds") or 0)
    return {
        "timesStartedAt": coverage.get("accountStartedAt"),
        "timeListened": listened,
        "timeSilenceRemoval": int(coverage.get("timeSilenceRemovalSeconds") or 0),
        "timeSkipping": int(coverage.get("timeSkippingSeconds") or 0),
        "timeIntroSkipping": int(coverage.get("timeIntroSkippingSeconds") or 0),
        "timeVariableSpeed": int(coverage.get("timeVariableSpeedSeconds") or 0),
    }


def repair_durations_offline() -> None:
    """Fix existing tables without calling the Pocket Casts API."""
    rows = load_episode_csv()
    # DictReader values are strings; coerce Year for aggregate grouping.
    for r in rows:
        try:
            r["Year"] = int(r["Year"])
        except (TypeError, ValueError):
            pass
    fixed = repair_missing_durations(rows)
    print(f"repaired durations from played-up-to on {fixed} rows")

    coverage_path = TABLES / "coverage.json"
    prior_coverage = (
        json.loads(coverage_path.read_text(encoding="utf-8"))
        if coverage_path.exists()
        else {}
    )
    year_counts_raw = prior_coverage.get("yearsProbed") or {
        str(y): 0 for y in range(2010, datetime.now(timezone.utc).year + 1)
    }
    stats = stats_from_coverage(prior_coverage)
    built, coverage = build_aggregates(rows, stats, year_counts_raw)
    # Keep prior account-level coverage fields; refresh listen bounds + sync time.
    for key in (
        "accountStartedAt",
        "timeListenedSeconds",
        "timeListenedHours",
        "timeSilenceRemovalSeconds",
        "timeSkippingSeconds",
        "timeIntroSkippingSeconds",
        "timeVariableSpeedSeconds",
        "yearsProbed",
        "note",
        "perEpisodeHistoryAvailableFrom",
    ):
        if key in prior_coverage:
            coverage[key] = prior_coverage[key]
    built["coverage"] = coverage
    built["aggregates"]["coverage"] = coverage
    built["meta"]["accountStartedAt"] = coverage.get("accountStartedAt")
    built["meta"]["accountTimeListenedHours"] = coverage.get("timeListenedHours")
    built["meta"]["perEpisodeHistoryFrom"] = coverage.get(
        "perEpisodeHistoryAvailableFrom"
    )
    built["meta"]["yearsWithoutHistory"] = [
        y for y, c in (coverage.get("yearsProbed") or {}).items() if c == 0
    ]

    write_tables(rows, built)
    write_excel(rows, coverage, built["meta"])
    sync_docs()
    run_standalone()
    print(
        f"duration repair complete: fixed={fixed} episodes={len(rows)} "
        f"totalDurationHours={built['meta'].get('totalDurationHours')}"
    )


def main() -> None:
    if "--repair-durations" in sys.argv:
        repair_durations_offline()
        return

    token = login()
    stats = fetch_stats(token)
    podcasts, folders = fetch_catalog(token)
    print("fetching history by year…")
    changes, year_counts_raw = fetch_history(token)
    if not changes:
        die("no history rows returned")
    print("fetching episode progress…")
    progress = fetch_progress(token, podcasts)
    starred = fetch_starred(token)

    old_compact: list[dict] = []
    eps_path = TABLES / "episodes.json"
    if eps_path.exists():
        old_compact = json.loads(eps_path.read_text(encoding="utf-8"))

    rows = build_rows(changes, podcasts, folders, progress, starred)
    print(f"built {len(rows)} episode rows; latest={rows[0]['Listened at (UTC)']}")

    built, coverage = build_aggregates(rows, stats, year_counts_raw)
    compact = write_tables(rows, built)
    # Transcripts are keyed by episode_idx. Daily re-exports reorder episodes, so
    # the sqlite store MUST be remapped by podcast+title before themes/notes use it.
    # notes.json is regenerated from the repaired store. Never leave sqlite poisoned
    # while only remapping notes — that caused cross-show theme contamination.
    sys.path.insert(0, str(ROOT / "scripts"))
    sqlite_path = ROOT / "transcripts" / "full.sqlite"
    sqlite_ready = False
    if sqlite_path.exists() and sqlite_path.stat().st_size > 1024:
        # Skip Git LFS pointer stubs (small text files starting with "version https://git-lfs").
        head = sqlite_path.read_bytes()[:64]
        sqlite_ready = not head.startswith(b"version https://git-lfs")
    if sqlite_ready:
        from transcript_store import (
            assert_aligned,
            connect,
            repair_episode_indices,
            write_slim_notes,
        )

        conn = connect()
        try:
            repair = repair_episode_indices(conn, compact)
            print(
                "transcript_index_repair "
                f"kept={repair['kept']} orphans_deleted={repair['orphansDeleted']} "
                f"collisions={repair['collisionsResolved']}"
            )
            align = assert_aligned(conn, compact)
            print(f"transcript_align_ok exact={align['exact']}")
            prior_notes = {}
            if (TABLES / "notes.json").exists():
                prior_notes = json.loads(
                    (TABLES / "notes.json").read_text(encoding="utf-8")
                )
            write_slim_notes(
                conn,
                TABLES / "notes.json",
                TABLES / "notes-meta.json",
                total_episodes=len(compact),
                preserve_notes=prior_notes,
                episodes=compact,
            )
        finally:
            conn.close()
    else:
        print(
            "transcript sqlite missing or Git LFS pointer — "
            "remapping notes.json only (run: git lfs pull)"
        )
        notes = remap_notes(old_compact, compact)
        (TABLES / "notes.json").write_text(
            json.dumps(notes, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    write_excel(rows, coverage, built["meta"])
    if os.environ.get("SKIP_THEMES") == "1":
        print("SKIP_THEMES=1 — leaving existing theme tables")
    else:
        run_themes()
    sync_docs()
    run_standalone()
    print(
        f"sync complete: episodes={len(rows)} latest={coverage.get('historyLatestListen')} "
        f"accountHours={coverage.get('timeListenedHours')}"
    )


if __name__ == "__main__":
    main()
